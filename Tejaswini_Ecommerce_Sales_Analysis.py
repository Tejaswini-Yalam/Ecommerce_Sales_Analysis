import os
import io
import random
from datetime import datetime, timedelta

import numpy as np
import pandas as pd
import streamlit as st
import plotly.express as px
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt


# =========================================================
# CONFIG 
# =========================================================

RANDOM_SEED = 42
N_ROWS = 600
DATA_DIR = "data"
RAW_PATH = os.path.join(DATA_DIR, "ecommerce_sales_raw.csv")
CLEAN_PATH = os.path.join(DATA_DIR, "ecommerce_sales_clean.csv")

REGIONS = ["North", "South", "East", "West", "Central"]
SEGMENTS = ["New", "Regular", "Premium"]
PAYMENT_METHODS = ["Credit Card", "UPI", "Net Banking", "COD", "Wallet"]
DELIVERY_STATUS = ["Delivered", "Delayed", "Cancelled"]

CATEGORY_PRODUCTS = {
    "Electronics": ["Wireless Earbuds", "Smartphone Case", "Bluetooth Speaker",
                    "Power Bank", "Smartwatch"],
    "Fashion": ["Cotton T-Shirt", "Denim Jeans", "Running Shoes", "Handbag", "Jacket"],
    "Home & Kitchen": ["Non-stick Pan", "Blender", "LED Lamp", "Storage Box", "Cushion Cover"],
    "Beauty": ["Face Serum", "Lipstick", "Shampoo", "Sunscreen", "Perfume"],
    "Sports": ["Yoga Mat", "Dumbbell Set", "Cricket Bat", "Football", "Cycling Gloves"],
    "Books": ["Fiction Novel", "Self-Help Book", "Cookbook", "Comic Book", "Biography"],
    "Grocery": ["Organic Rice 5kg", "Green Tea Pack", "Almonds 500g", "Olive Oil 1L", "Protein Bar Pack"],
}

BASE_PRICE = {
    "Electronics": (800, 6000), "Fashion": (300, 3000), "Home & Kitchen": (250, 3500),
    "Beauty": (150, 1800), "Sports": (300, 4000), "Books": (150, 900), "Grocery": (100, 1200),
}


# =========================================================
# STEP 1: DATA GENERATION  (creates a deliberately messy raw dataset)
# =========================================================
def _random_date(start, end):
    return start + timedelta(days=random.randint(0, (end - start).days))


def _make_row(order_id):
    category = random.choice(list(CATEGORY_PRODUCTS.keys()))
    product = random.choice(CATEGORY_PRODUCTS[category])
    low, high = BASE_PRICE[category]
    unit_price = round(random.uniform(low, high), 2)
    quantity = random.randint(1, 6)
    discount_percent = random.choice([0, 0, 5, 10, 15, 20, 25])
    order_date = _random_date(datetime(2024, 1, 1), datetime(2024, 12, 31))
    delivery_days = random.randint(1, 10)
    rating = random.choice([1, 2, 3, 4, 4, 5, 5, 5, np.nan])

    # intentional data-quality issues, cleaned in step 2
    if random.random() < 0.03:
        unit_price = -unit_price
    if random.random() < 0.03:
        quantity = 0
    if random.random() < 0.04:
        delivery_days = np.nan
    if random.random() < 0.05:
        category = category.upper()
    if random.random() < 0.02:
        product = "  " + product + "  "

    total_sales = round(unit_price * quantity * (1 - discount_percent / 100), 2)

    return {
        "order_id": f"ORD-{order_id:05d}",
        "order_date": order_date.strftime("%Y-%m-%d"),
        "customer_id": f"CUST-{random.randint(1000, 1250)}",
        "customer_segment": random.choice(SEGMENTS),
        "region": random.choice(REGIONS),
        "product_category": category,
        "product_name": product,
        "unit_price": unit_price,
        "quantity": quantity,
        "discount_percent": discount_percent,
        "payment_method": random.choice(PAYMENT_METHODS),
        "delivery_days": delivery_days,
        "delivery_status": random.choice(DELIVERY_STATUS),
        "customer_rating": rating,
        "total_sales": total_sales,
    }


@st.cache_data
def generate_raw_data() -> pd.DataFrame:
    random.seed(RANDOM_SEED)
    np.random.seed(RANDOM_SEED)
    rows = [_make_row(i + 1) for i in range(N_ROWS)]
    df = pd.DataFrame(rows)
    dup_sample = df.sample(8, random_state=RANDOM_SEED)
    df = pd.concat([df, dup_sample], ignore_index=True)
    return df


# =========================================================
# STEP 2: DATA CLEANING
# =========================================================
@st.cache_data
def clean_data(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()

    before = len(df)
    df = df.drop_duplicates()
    removed_dupes = before - len(df)

    text_cols = ["product_category", "product_name", "region",
                 "customer_segment", "payment_method", "delivery_status"]
    for col in text_cols:
        df[col] = df[col].astype(str).str.strip()
    df["product_category"] = df["product_category"].str.title()

    df["unit_price"] = df["unit_price"].abs()
    df.loc[df["quantity"] <= 0, "quantity"] = 1

    df["delivery_days"] = df.groupby("product_category")["delivery_days"] \
        .transform(lambda s: s.fillna(s.median()))
    df["delivery_days"] = df["delivery_days"].fillna(df["delivery_days"].median())

    df["rating_missing"] = df["customer_rating"].isna()

    df["total_sales"] = (
        df["unit_price"] * df["quantity"] * (1 - df["discount_percent"] / 100)
    ).round(2)

    df["order_date"] = pd.to_datetime(df["order_date"], errors="coerce")

    df.attrs["cleaning_report"] = {
        "rows_after_dedup": len(df),
        "duplicate_rows_removed": removed_dupes,
        "missing_ratings": int(df["rating_missing"].values.sum()),
    }
    return df


# =========================================================
# STEP 3: ANALYSIS FUNCTIONS
# =========================================================
def kpi_summary(df: pd.DataFrame) -> dict:
    avg_rating = df["customer_rating"].mean()
    return {
        "total_orders": int(len(df)),
        "total_revenue": round(df["total_sales"].sum(), 2),
        "avg_order_value": round(df["total_sales"].mean(), 2),
        "avg_rating": round(avg_rating, 2) if pd.notna(avg_rating) else None,
        "cancelled_rate_pct": round((df["delivery_status"] == "Cancelled").mean() * 100, 2),
    }


def revenue_by_category(df):
    return (df.groupby("product_category")["total_sales"].sum()
            .sort_values(ascending=False).reset_index()
            .rename(columns={"total_sales": "revenue"}))


def revenue_by_region(df):
    return (df.groupby("region")["total_sales"].sum()
            .sort_values(ascending=False).reset_index()
            .rename(columns={"total_sales": "revenue"}))


def revenue_by_segment(df):
    return (df.groupby("customer_segment")["total_sales"]
            .agg(revenue="sum", orders="count", avg_order_value="mean")
            .reset_index().sort_values("revenue", ascending=False))


def monthly_revenue_trend(df):
    tmp = df.copy()
    tmp["month"] = tmp["order_date"].dt.to_period("M").astype(str)
    return tmp.groupby("month")["total_sales"].sum().reset_index()


def delivery_performance(df):
    return (df.groupby("delivery_status")
            .agg(orders=("order_id", "count"), avg_delivery_days=("delivery_days", "mean"))
            .reset_index())


def top_products(df, n=10):
    return (df.groupby("product_name")["total_sales"].sum()
            .sort_values(ascending=False).head(n).reset_index()
            .rename(columns={"total_sales": "revenue"}))


# =========================================================
# STEP 4: PPTX REPORT BUILDER (in-memory, no temp files needed)
# =========================================================
def _fig_to_buf(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


def build_pptx_report(df: pd.DataFrame) -> io.BytesIO:
    kpis = kpi_summary(df)

    cat_df = revenue_by_category(df)
    fig, ax = plt.subplots(figsize=(7, 4))
    ax.bar(cat_df["product_category"], cat_df["revenue"], color="#2E86AB")
    ax.set_title("Revenue by Product Category")
    plt.xticks(rotation=30, ha="right")
    cat_buf = _fig_to_buf(fig)

    reg_df = revenue_by_region(df)
    fig, ax = plt.subplots(figsize=(6, 6))
    ax.pie(reg_df["revenue"], labels=reg_df["region"], autopct="%1.1f%%")
    ax.set_title("Revenue Share by Region")
    reg_buf = _fig_to_buf(fig)

    trend_df = monthly_revenue_trend(df)
    fig, ax = plt.subplots(figsize=(7, 4))
    ax.plot(trend_df["month"], trend_df["total_sales"], marker="o", color="#A23B72")
    ax.set_title("Monthly Revenue Trend")
    plt.xticks(rotation=45, ha="right")
    trend_buf = _fig_to_buf(fig)

    prs = Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "E-Commerce Sales & Customer Behavior Analysis"
    slide.placeholders[1].text = (
        f"Total Revenue: Rs {kpis['total_revenue']:,.0f}  |  "
        f"Orders: {kpis['total_orders']:,}  |  Avg Rating: {kpis['avg_rating']}"
    )

    for title, img_buf in [
        ("Revenue by Product Category", cat_buf),
        ("Revenue Share by Region", reg_buf),
        ("Monthly Revenue Trend", trend_buf),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        slide.shapes.add_picture(img_buf, Inches(0.7), Inches(1.3), width=Inches(8.5))

    top_df = top_products(df, 10)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Top 10 Products by Revenue"
    rows, cols = len(top_df) + 1, top_df.shape[1]
    table = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(1.3),
                                    Inches(9), Inches(0.4 * rows)).table
    for j, col_name in enumerate(top_df.columns):
        table.cell(0, j).text = str(col_name)
        table.cell(0, j).text_frame.paragraphs[0].font.bold = True
    for i, (_, row) in enumerate(top_df.iterrows(), start=1):
        for j, val in enumerate(row):
            table.cell(i, j).text = f"{val:,.2f}" if isinstance(val, float) else str(val)

    out_buf = io.BytesIO()
    prs.save(out_buf)
    out_buf.seek(0)
    return out_buf


# =========================================================
# STEP 5: STREAMLIT DASHBOARD (main app)
# =========================================================
def main():
    st.set_page_config(page_title="E-Commerce Sales Analysis", layout="wide")
    st.title("🛒 E-Commerce Sales & Customer Behavior Dashboard")
    st.caption("Analyze revenue, regions, categories, delivery performance and customer segments.")

    # --- build pipeline (also saved to disk so you can inspect the CSVs) ---
    raw_df = generate_raw_data()
    df = clean_data(raw_df)

    os.makedirs(DATA_DIR, exist_ok=True)
    if not os.path.exists(RAW_PATH):
        raw_df.to_csv(RAW_PATH, index=False)
    if not os.path.exists(CLEAN_PATH):
        df.to_csv(CLEAN_PATH, index=False)

    # --- sidebar filters ---
    st.sidebar.header("Filters")
    regions = st.sidebar.multiselect("Region", sorted(df["region"].unique()),
                                      default=sorted(df["region"].unique()))
    categories = st.sidebar.multiselect("Product Category", sorted(df["product_category"].unique()),
                                         default=sorted(df["product_category"].unique()))
    segments = st.sidebar.multiselect("Customer Segment", sorted(df["customer_segment"].unique()),
                                       default=sorted(df["customer_segment"].unique()))
    date_min, date_max = df["order_date"].min(), df["order_date"].max()
    date_range = st.sidebar.date_input("Order date range", (date_min, date_max))

    mask = (df["region"].isin(regions) & df["product_category"].isin(categories)
            & df["customer_segment"].isin(segments))
    if isinstance(date_range, tuple) and len(date_range) == 2:
        start, end = pd.to_datetime(date_range[0]), pd.to_datetime(date_range[1])
        mask &= df["order_date"].between(start, end)
    filtered = df[mask]

    if filtered.empty:
        st.warning("No data matches the selected filters.")
        st.stop()

    # --- KPIs ---
    kpis = kpi_summary(filtered)
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Total Orders", f"{kpis['total_orders']:,}")
    c2.metric("Total Revenue", f"Rs {kpis['total_revenue']:,.0f}")
    c3.metric("Avg Order Value", f"Rs {kpis['avg_order_value']:,.0f}")
    c4.metric("Avg Rating", f"{kpis['avg_rating']:.2f}" if kpis["avg_rating"] else "N/A")
    c5.metric("Cancelled Rate", f"{kpis['cancelled_rate_pct']}%")
    st.divider()

    # --- charts ---
    col1, col2 = st.columns(2)
    with col1:
        st.subheader("Revenue by Product Category")
        fig = px.bar(revenue_by_category(filtered), x="product_category", y="revenue",
                     color="product_category", text_auto=".2s")
        fig.update_layout(showlegend=False, xaxis_title="", yaxis_title="Revenue")
        st.plotly_chart(fig, width="stretch")

    with col2:
        st.subheader("Revenue by Region")
        fig = px.pie(revenue_by_region(filtered), names="region", values="revenue", hole=0.4)
        st.plotly_chart(fig, width="stretch")

    col3, col4 = st.columns(2)
    with col3:
        st.subheader("Monthly Revenue Trend")
        fig = px.line(monthly_revenue_trend(filtered), x="month", y="total_sales", markers=True)
        fig.update_layout(xaxis_title="Month", yaxis_title="Revenue")
        st.plotly_chart(fig, width="stretch")

    with col4:
        st.subheader("Customer Segment Performance")
        fig = px.bar(revenue_by_segment(filtered), x="customer_segment", y="revenue",
                     color="customer_segment", text_auto=".2s")
        fig.update_layout(showlegend=False, xaxis_title="", yaxis_title="Revenue")
        st.plotly_chart(fig, width="stretch")

    st.divider()
    col5, col6 = st.columns(2)
    with col5:
        st.subheader("Top 10 Products by Revenue")
        st.dataframe(top_products(filtered, 10), width="stretch", hide_index=True)
    with col6:
        st.subheader("Delivery Performance")
        st.dataframe(delivery_performance(filtered), width="stretch", hide_index=True)

    st.divider()
    st.subheader("Filtered Raw Data")
    st.dataframe(filtered, width="stretch", hide_index=True)

    # --- report download ---
    st.divider()
    st.subheader("📊 Download PowerPoint Report")
    if st.button("Generate Report"):
        with st.spinner("Building PPTX report..."):
            pptx_buf = build_pptx_report(filtered)
        st.download_button(
            "⬇️ Download Ecommerce_Sales_Report.pptx",
            data=pptx_buf,
            file_name="Ecommerce_Sales_Report.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )


if __name__ == "__main__":
    main()
      